"""FlexSizer - Cisco Flex 3.0 EA licence sizing; Tkinter front end for flex_ea_calc.

Copyright (c) 2026 Nithin Eluvathingal. MIT License - see LICENSE.
"""

import copy
import json
import os
import tkinter as tk
import webbrowser
from tkinter import filedialog, messagebox, ttk

from flex_ea_calc import (DATA_SHEET_URL, DEFAULT_SETTINGS, LICENSE_TYPES, SKUS, TREATMENTS,
                          calculate, kw_scenarios, marginal_costs, summary_text, unit_prices)

APP_NAME = "FlexSizer"
APP_TAGLINE = "Cisco Collaboration Flex 3.0 EA sizing"
_APPDATA = os.environ.get("APPDATA", os.path.expanduser("~"))
SETTINGS_DIR = os.path.join(_APPDATA, APP_NAME)
SETTINGS_FILE = os.path.join(SETTINGS_DIR, "settings.json")
OLD_SETTINGS_FILE = os.path.join(_APPDATA, "Cisco Flex EA Calculator", "settings.json")  # before the rename
SETTINGS_VERSION = 4   # v2: data sheet rules; v3: pricing, 20% growth; v4: Cisco user-type mapping. Older saves are ignored
TREATMENT_BY_LABEL = {v: k for k, v in TREATMENTS.items()}
RULE_FIELDS = [   # (settings key, label, number type)
    ("ca_pct", "Common Area included (% of KW)", float),
    ("access_pct", "Access included (% of KW)", float),
    ("min_ea", "EA minimum (KW)", int),
    ("growth_allowance_pct", "Growth allowance (%)", float),
    ("premises_reg_pct", "Premises device registration (% of KW)", float),
    ("headroom_pct", "Planning head-room on quantities (%)", float),
]
PRICE_FIELDS = [
    ("price_kw", f"EA Knowledge Worker  {SKUS['kw']}", float),
    ("price_ca", f"Common Area add-on  {SKUS['ca']}", float),
    ("price_access", f"Access add-on  {SKUS['access']}", float),
    ("discount_pct", "Discount off list (%)", float),
    ("term_months", "Term (months)", int),
]
WHATIF_FIELDS = [("kw", "+ Knowledge workers"), ("ca", "+ Common Area"), ("access", "+ Access")]

C = {"bg": "#F3F6FA", "surface": "#FFFFFF", "ink": "#1E2A38", "muted": "#6B7785", "line": "#D6DDE6",
     "stripe": "#F7F9FC", "accent": "#049FD9", "accent_dark": "#0379A6", "accent_soft": "#DDF2FB",
     "navy": "#0B2A4A", "free": "#3BB273", "orange": "#E07B39", "track": "#E9EEF4"}
# One colour per Flex user type, used for the row dots, tiles and bars.
TYPE_COLOR = {"KW": "#049FD9", "CA": "#14A38B", "ACCESS": "#7B5CD6", "FREE": "#A0AAB6"}
F = "Segoe UI"


def money(v):
    return f"{v:,.2f}"


def load_settings():
    s = copy.deepcopy(DEFAULT_SETTINGS)
    path = SETTINGS_FILE if os.path.exists(SETTINGS_FILE) else OLD_SETTINGS_FILE
    try:
        with open(path, encoding="utf-8") as f:
            saved = json.load(f)
        if saved.get("version") == SETTINGS_VERSION:
            treatment = saved.pop("treatment", {})
            s["treatment"].update({k: v for k, v in treatment.items() if v in TREATMENTS})
            s.update({k: v for k, v in saved.items() if k in s})
    except (OSError, ValueError):
        pass
    return s


def save_settings(s):
    os.makedirs(SETTINGS_DIR, exist_ok=True)
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump({**s, "version": SETTINGS_VERSION}, f, indent=2)


def to_number(text, kind=int):
    try:
        return max(0, kind(str(text).strip() or 0))
    except ValueError:
        return 0


class Tile(tk.Frame):
    """KPI card: coloured top stripe, big number, caption and a small detail line."""

    def __init__(self, master, caption, color):
        super().__init__(master, bg=C["surface"], highlightbackground=C["line"], highlightthickness=1)
        tk.Frame(self, bg=color, height=4).pack(fill="x")
        body = tk.Frame(self, bg=C["surface"], padx=12, pady=8)
        body.pack(fill="both", expand=True)
        self.value = tk.Label(body, text="0", bg=C["surface"], fg=color, font=(F, 18, "bold"), anchor="w")
        self.value.pack(fill="x")
        cap = tk.Label(body, text=caption, bg=C["surface"], fg=C["ink"], font=(F, 9, "bold"), anchor="w",
                       justify="left")
        cap.pack(fill="x")
        self.detail = tk.Label(body, text="", bg=C["surface"], fg=C["muted"], font=(F, 8), anchor="w",
                               justify="left")
        self.detail.pack(fill="x")
        for lbl in (cap, self.detail):   # wrap instead of clipping on narrow windows
            lbl.bind("<Configure>", lambda e: e.widget.config(wraplength=e.width - 6))

    def set(self, value, detail=""):
        self.value.config(text=value)
        self.detail.config(text=detail)


class App(ttk.Frame):
    def __init__(self, master):
        super().__init__(master, padding=14, style="App.TFrame")
        self.scale = master.winfo_fpixels("1i") / 96.0   # pixel sizes below are for 100% Windows scaling
        self.settings = load_settings()
        self.qty_vars = {lic: tk.StringVar(value="0") for lic in LICENSE_TYPES}
        self.how_vars = {lic: tk.StringVar(value=TREATMENTS[self.settings["treatment"][lic]])
                         for lic in LICENSE_TYPES}
        self.rule_vars = {key: tk.StringVar(value=f"{self.settings[key]:g}")
                          for key, _, _ in RULE_FIELDS + PRICE_FIELDS}
        self.excess_mode = tk.StringVar(value=self.settings["excess_mode"])
        self.whatif_vars = {key: tk.StringVar(value="0") for key, _ in WHATIF_FIELDS}
        self.result = None
        self._build()
        for v in [*self.qty_vars.values(), *self.how_vars.values(), *self.rule_vars.values(),
                  *self.whatif_vars.values(), self.excess_mode]:
            v.trace_add("write", lambda *_: self.recalc())
        self.recalc()

    def px(self, n):
        return int(n * self.scale)

    # ---------------------------------------------------------------- layout
    def _build(self):
        self.columnconfigure(1, weight=1)
        self.rowconfigure(0, weight=1)
        left = ttk.Frame(self, style="App.TFrame")
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 14))
        left.rowconfigure(1, weight=1)
        right = ttk.Frame(self, style="App.TFrame")
        right.grid(row=0, column=1, sticky="nsew")
        right.columnconfigure(0, weight=1)
        right.rowconfigure(3, weight=1)
        self._build_inputs(left)
        self._build_rules(left)
        foot = ttk.Frame(left, style="App.TFrame")
        foot.grid(row=2, column=0, sticky="ew", pady=(8, 0))
        ttk.Label(foot, text="© 2026 Nithin Eluvathingal · MIT License", style="Foot.TLabel").pack(side="left")
        link = ttk.Label(foot, text="Cisco Flex 3.0 data sheet ↗", style="FootLink.TLabel", cursor="hand2")
        link.pack(side="right")
        link.bind("<Button-1>", lambda _: webbrowser.open(DATA_SHEET_URL))
        self._build_tiles(right)
        self._build_charts(right)
        self._build_tables(right)

    def _build_inputs(self, parent):
        inp = ttk.LabelFrame(parent, text="1  Existing perpetual licences")
        inp.grid(row=0, column=0, sticky="nsew")
        for col, text in enumerate(("Licence type", "Quantity", "Flex user type")):
            ttk.Label(inp, text=text, style="ColHead.TLabel").grid(row=0, column=col, sticky="w",
                                                                   padx=(0 if col == 0 else 8, 0), pady=(0, 4))
        self.dots = {}
        for i, lic in enumerate(LICENSE_TYPES, start=1):
            cell = ttk.Frame(inp)
            cell.grid(row=i, column=0, sticky="w", pady=3)
            self.dots[lic] = tk.Label(cell, text="●", bg=C["surface"], font=(F, 11))
            self.dots[lic].pack(side="left", padx=(0, 6))
            ttk.Label(cell, text=lic).pack(side="left")
            e = ttk.Spinbox(inp, from_=0, to=1_000_000, width=9, textvariable=self.qty_vars[lic])
            e.grid(row=i, column=1, padx=8, pady=3)
            if i == 1:
                self.first_entry = e
            ttk.Combobox(inp, values=list(TREATMENTS.values()), state="readonly", width=33,
                         textvariable=self.how_vars[lic]).grid(row=i, column=2, pady=3)
        ttk.Button(inp, text="Clear quantities", command=self.clear).grid(
            row=len(LICENSE_TYPES) + 1, column=0, columnspan=3, sticky="w", pady=(10, 0))

    def _build_rules(self, parent):
        box = ttk.LabelFrame(parent, text="2  Rules & pricing")
        box.grid(row=1, column=0, sticky="nsew", pady=(14, 0))
        box.columnconfigure(0, weight=1)
        box.rowconfigure(0, weight=1)
        tabs = ttk.Notebook(box)
        tabs.grid(row=0, column=0, sticky="nsew")

        rules = ttk.Frame(tabs, padding=(4, 8))
        for i, (key, label, _) in enumerate(RULE_FIELDS):
            ttk.Label(rules, text=label).grid(row=i, column=0, sticky="w", pady=2)
            ttk.Entry(rules, width=9, textvariable=self.rule_vars[key], justify="right").grid(
                row=i, column=1, sticky="w", padx=8)
        row = len(RULE_FIELDS)
        ttk.Label(rules, text="Devices above the free allowance:").grid(
            row=row, column=0, columnspan=2, sticky="w", pady=(8, 0))
        rb = ttk.Frame(rules)
        rb.grid(row=row + 1, column=0, columnspan=2, sticky="w")
        ttk.Radiobutton(rb, text="Add-on licences", value="addon", variable=self.excess_mode).pack(side="left")
        ttk.Radiobutton(rb, text="Extra EA KW licences", value="ea",
                        variable=self.excess_mode).pack(side="left", padx=(16, 0))
        tabs.add(rules, text="EA rules")

        price = ttk.Frame(tabs, padding=(4, 8))
        for i, (key, label, _) in enumerate(PRICE_FIELDS):
            ttk.Label(price, text=label).grid(row=i, column=0, sticky="w", pady=2)
            ttk.Entry(price, width=9, textvariable=self.rule_vars[key], justify="right").grid(
                row=i, column=1, sticky="w", padx=8)
        ttk.Label(price, text="USD list price per licence per month (CCW).", style="Muted.TLabel").grid(
            row=len(PRICE_FIELDS), column=0, columnspan=2, sticky="w", pady=(10, 0))
        tabs.add(price, text="Pricing")

        btns = ttk.Frame(box)
        btns.grid(row=1, column=0, sticky="w", pady=(10, 0))
        ttk.Button(btns, text="Save as default", command=self.save_defaults).pack(side="left")
        ttk.Button(btns, text="Reset to data sheet", command=self.reset_rules).pack(side="left", padx=8)

    def _build_tiles(self, parent):
        row = ttk.Frame(parent, style="App.TFrame")
        row.grid(row=0, column=0, sticky="ew")
        specs = [("ea", "EA KW licences", TYPE_COLOR["KW"]),
                 ("ca", "Common Area add-ons", TYPE_COLOR["CA"]),
                 ("access", "Access add-ons", TYPE_COLOR["ACCESS"]),
                 ("monthly", "Per month (USD)", C["navy"]),
                 ("term", "Term total (USD)", C["orange"])]
        self.tiles = {}
        for i, (key, caption, color) in enumerate(specs):
            row.columnconfigure(i, weight=1, uniform="tile")
            self.tiles[key] = Tile(row, caption, color)
            self.tiles[key].grid(row=0, column=i, sticky="nsew", padx=(0 if i == 0 else 10, 0))
        self.ea_label = self.tiles["ea"].value
        self.note_label = tk.Label(parent, text="", bg=C["accent_soft"], fg=C["navy"], font=(F, 9),
                                   anchor="w", justify="left", padx=12, pady=7)
        self.note_label.grid(row=1, column=0, sticky="ew", pady=(12, 0))
        self.note_label.bind("<Configure>", lambda e: self.note_label.config(wraplength=e.width - 24))

    def _build_charts(self, parent):
        card = tk.Frame(parent, bg=C["surface"], highlightbackground=C["line"], highlightthickness=1)
        card.grid(row=2, column=0, sticky="ew", pady=(12, 0))
        self.chart = tk.Canvas(card, height=self.px(150), bg=C["surface"], highlightthickness=0)
        self.chart.pack(fill="x", padx=14, pady=10)
        self.chart.bind("<Configure>", lambda e: self.draw_chart())

    def _build_tables(self, parent):
        tabs = ttk.Notebook(parent)
        tabs.grid(row=3, column=0, sticky="nsew", pady=(12, 0))

        summ = ttk.Frame(tabs, padding=8)
        self.tree = self._table(summ, ("value",), ("",), (130,), first=("Item", 340), headings=False)
        tabs.add(summ, text="Summary & cost")

        sc = ttk.Frame(tabs, padding=8)
        self.scen = self._table(sc, ("kw", "ca", "access", "monthly", "term", "delta"),
                                ("EA KW", "CA add-on", "Access add-on", "Monthly", "Term total", "Extra cost"),
                                (70, 85, 110, 100, 115, 100), first=("Option", 180))
        self.scen_note = ttk.Label(sc, text="", style="Muted.TLabel", justify="left")
        self.scen_note.pack(anchor="w", fill="x", pady=(8, 0))
        self.scen_note.bind("<Configure>", lambda e: self.scen_note.config(wraplength=e.width))
        tabs.add(sc, text="More KW vs add-ons")

        wi = ttk.Frame(tabs, padding=8)
        inputs = ttk.Frame(wi)
        inputs.pack(fill="x")
        for i, (key, label) in enumerate(WHATIF_FIELDS):
            ttk.Label(inputs, text=label).grid(row=0, column=i * 2, sticky="w", padx=(0 if i == 0 else 16, 6))
            ttk.Spinbox(inputs, from_=0, to=1_000_000, width=8, textvariable=self.whatif_vars[key]).grid(
                row=0, column=i * 2 + 1, sticky="w")
        self.whatif = self._table(wi, ("now", "after", "delta"), ("Current", "After change", "Difference"),
                                  (130, 130, 130), first=("", 220), pady=(10, 0))
        self.unit_note = ttk.Label(wi, text="", style="Muted.TLabel", justify="left")
        self.unit_note.pack(anchor="w", fill="x", pady=(8, 0))
        self.unit_note.bind("<Configure>", lambda e: self.unit_note.config(wraplength=e.width))
        tabs.add(wi, text="What-if: add devices")

    def _table(self, parent, cols, heads, widths, first, headings=True, pady=0):
        box = ttk.Frame(parent)
        box.pack(fill="both", expand=True, pady=pady)
        tree = ttk.Treeview(box, columns=cols, show="tree headings" if headings else "tree", height=9)
        tree.heading("#0", text=first[0])
        for c, h in zip(cols, heads):
            tree.heading(c, text=h)
            tree.column(c, anchor="e", stretch=False)
        # Widths are proportions: share the table's actual width so no column is clipped.
        shares = list(zip(("#0",) + tuple(cols), (first[1],) + tuple(widths)))
        total = sum(w for _, w in shares)

        def fit(event):
            for c, w in shares:
                tree.column(c, width=max(self.px(40), int(event.width * w / total)))
        tree.bind("<Configure>", fit)
        tree.tag_configure("odd", background=C["stripe"])
        tree.tag_configure("key", font=(F, 10, "bold"), foreground=C["navy"])
        tree.tag_configure("section", font=(F, 10, "bold"))
        tree.tag_configure("up", foreground=C["orange"])
        bar = ttk.Scrollbar(box, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=bar.set)
        tree.pack(side="left", fill="both", expand=True)
        bar.pack(side="right", fill="y")
        return tree

    def _rows(self, tree, rows, parent=""):
        """Insert (text, values, extra tags) rows with zebra striping."""
        start = len(tree.get_children(parent))
        ids = []
        for i, (text, values, tags) in enumerate(rows, start=start):
            ids.append(tree.insert(parent, "end", text=text, values=values,
                                   tags=(("odd",) if i % 2 else ()) + tuple(tags), open=True))
        return ids

    # ---------------------------------------------------------------- logic
    def current_settings(self):
        s = {key: to_number(self.rule_vars[key].get(), kind) for key, _, kind in RULE_FIELDS + PRICE_FIELDS}
        s["treatment"] = {lic: TREATMENT_BY_LABEL[self.how_vars[lic].get()] for lic in LICENSE_TYPES}
        s["excess_mode"] = self.excess_mode.get()
        return s

    def result_rows(self, r, s):
        return [
            ("Professional (Knowledge Workers)", r.knowledge_workers),
            ("Flex EA Knowledge Worker licences", r.ea_licenses),
            ("Enhanced (Task Worker / Common Area)", r.ca_devices),
            (f"   Included ({s['ca_pct']:g}% of {r.ea_licenses:,})", r.ca_included),
            ("   Common Area add-on licences", r.ca_addon),
            ("   Unused Common Area allowance", r.ca_unused),
            ("Access (Analog)", r.access_devices),
            (f"   Included ({s['access_pct']:g}% of {r.ea_licenses:,})", r.access_included),
            ("   Access add-on licences", r.access_addon),
            ("   Unused Access allowance", r.access_unused),
            ("Not counted", r.not_counted),
            (f"Growth allowance ceiling (+{s['growth_allowance_pct']:g}%)", r.growth_ceiling),
            (f"Premises device registration ({s['premises_reg_pct']:g}%)", r.premises_reg),
        ]

    def cost_rows(self, r, s):
        term = max(1, s["term_months"])
        return [
            (f"{SKUS['kw']} monthly", r.cost_kw),
            (f"{SKUS['ca']} monthly", r.cost_ca),
            (f"{SKUS['access']} monthly", r.cost_access),
            ("Total per month", r.monthly),
            ("Total per year", r.monthly * 12),
            (f"Total for {term}-month term", r.term_total),
        ]

    def quantities(self, extra=None):
        q = {lic: to_number(v.get()) for lic, v in self.qty_vars.items()}
        if extra:
            # Put what-if extras on the first licence type counted in that bucket.
            treatment = self.current_settings()["treatment"]
            for bucket, add in (("KW", extra["kw"]), ("CA", extra["ca"]), ("ACCESS", extra["access"])):
                lic = next((l for l in LICENSE_TYPES if treatment[l] == bucket), None)
                if lic and add:
                    q[lic] += add
        return q

    def recalc(self):
        quantities = self.quantities()
        s = self.current_settings()
        r = self.result = calculate(quantities, s)
        term = max(1, s["term_months"])

        for lic in LICENSE_TYPES:
            self.dots[lic].config(fg=TYPE_COLOR[s["treatment"][lic]])
        self.tiles["ea"].set(f"{r.ea_licenses:,}", f"{r.knowledge_workers:,} knowledge workers"
                             + (" · EA minimum" if r.min_applied else ""))
        self.tiles["ca"].set(f"{r.ca_addon:,}", f"{r.ca_included:,} included free")
        self.tiles["access"].set(f"{r.access_addon:,}", f"{r.access_included:,} included free")
        self.tiles["monthly"].set(f"{r.monthly:,.0f}", f"{r.monthly * 12:,.0f} per year")
        self.tiles["term"].set(f"{r.term_total:,.0f}", f"{term} months"
                               + (f" · {s['discount_pct']:g}% off list" if s["discount_pct"] else " · list price"))

        notes = []
        if r.min_applied:
            notes.append(f"Raised to the EA minimum of {s['min_ea']:,} KWs "
                         f"({r.spare_kw:,} spare also cover devices above the allowance).")
        elif r.spare_kw:
            notes.append(f"Includes {r.spare_kw:,} extra KW licences for devices above the included allowance.")
        notes.append(f"You can grow to {r.growth_ceiling:,} KWs inside the "
                     f"{s['growth_allowance_pct']:g}% growth allowance.")
        self.note_label.config(text="ⓘ  " + " ".join(notes))

        self.tree.delete(*self.tree.get_children())
        key_rows = {"Flex EA Knowledge Worker licences", "   Common Area add-on licences", "   Access add-on licences"}
        self._rows(self.tree, [(label, (f"{value:,}",), ("key",) if label in key_rows else ())
                               for label, value in self.result_rows(r, s)])
        (cost,) = self._rows(self.tree, [(f"Cost (USD, {s['discount_pct']:g}% discount)", ("",), ("section",))])
        self._rows(self.tree, [(label, (money(value),), ("key",) if label.startswith("Total for") else ())
                               for label, value in self.cost_rows(r, s)], cost)
        if s["headroom_pct"]:
            (parent,) = self._rows(self.tree, [(f"Quantities after {s['headroom_pct']:g}% head-room", ("",),
                                                ("section",))])
            self._rows(self.tree, [(f"{lic} ({qty:,})", (f"{sized:,}",), ()) for lic, qty, sized, _ in r.lines],
                       parent)
        self.fill_scenarios(quantities, s)
        self.fill_whatif(s)
        self.draw_chart()

    def draw_chart(self):
        r, cv = self.result, self.chart
        if r is None:
            return
        cv.delete("all")
        w = max(cv.winfo_width(), self.px(300))
        bar_h, gap = self.px(16), self.px(8)
        title = (F, 9, "bold")
        small = (F, 8)

        # Monthly cost split by SKU.
        cv.create_text(0, 0, text="Monthly cost split", anchor="nw", font=title, fill=C["ink"])
        y = self.px(20)
        parts = [("EA KW", r.cost_kw, TYPE_COLOR["KW"]), ("Common Area", r.cost_ca, TYPE_COLOR["CA"]),
                 ("Access", r.cost_access, TYPE_COLOR["ACCESS"])]
        total = sum(p[1] for p in parts)
        cv.create_rectangle(0, y, w, y + bar_h, fill=C["track"], width=0)
        x = 0
        for _, value, color in parts:
            if total and value:
                seg = w * value / total
                cv.create_rectangle(x, y, x + seg, y + bar_h, fill=color, width=0)
                x += seg
        lx, ly = 0, y + bar_h + gap
        for name, value, color in parts:
            share = f"{value / total:.0%}" if total else "0%"
            cv.create_rectangle(lx, ly + 2, lx + self.px(10), ly + self.px(12), fill=color, width=0)
            t = cv.create_text(lx + self.px(14), ly, anchor="nw", font=small, fill=C["muted"],
                               text=f"{name}  {money(value)}  ({share})")
            lx = cv.bbox(t)[2] + self.px(18)

        # Free allowance use: green = covered free, colour = add-on (or spare KW), grey = unused.
        y = ly + self.px(28)
        cv.create_text(0, y, text="Free allowance use", anchor="nw", font=title, fill=C["ink"])
        label_w = self.px(110)
        for name, devices, included, addon, color in (
                ("Common Area", r.ca_devices, r.ca_included, r.ca_addon, TYPE_COLOR["CA"]),
                ("Access", r.access_devices, r.access_included, r.access_addon, TYPE_COLOR["ACCESS"])):
            y += bar_h + gap
            cv.create_text(0, y + bar_h / 2, text=name, anchor="w", font=small, fill=C["ink"])
            x0, x1 = label_w, w - self.px(230)
            span = max(devices, included, 1)
            cv.create_rectangle(x0, y, x1, y + bar_h, fill=C["track"], width=0)
            free = min(devices, included)
            fx = x0 + (x1 - x0) * free / span
            cv.create_rectangle(x0, y, fx, y + bar_h, fill=C["free"], width=0)
            cv.create_rectangle(fx, y, x0 + (x1 - x0) * devices / span, y + bar_h, fill=color, width=0)
            other = devices - free - addon   # covered by spare EA licences
            cv.create_text(x1 + self.px(10), y + bar_h / 2, anchor="w", font=small, fill=C["muted"],
                           text=f"{free:,} free · {addon:,} add-on"
                                + (f" · {other:,} on spare KW" if other > 0 else "")
                                + (f" · {included - devices:,} unused" if included > devices else ""))

    def fill_scenarios(self, quantities, s):
        self.scen.delete(*self.scen.get_children())
        options = kw_scenarios(quantities, s)
        base = options[0][1].term_total
        self._rows(self.scen, [(
            label,
            (f"{o.ea_licenses:,}", f"{o.ca_addon:,}", f"{o.access_addon:,}", money(o.monthly),
             money(o.term_total), ("+" if o.term_total > base else "") + money(o.term_total - base)),
            ("key",) if i == 0 else ("up",) if o.term_total > base else ())
            for i, (label, o) in enumerate(options)])
        m = marginal_costs(s)
        self.scen_note.config(text=(
            f"Each extra EA KW costs {money(m['kw'])}/month but only frees {s['ca_pct']:g}% of a Common Area "
            f"and {s['access_pct']:g}% of an Access add-on, so it adds a net {money(m['kw_net'])}/month. "
            + ("Buying KWs only for your real knowledge workers (or the EA minimum) and add-ons for the rest "
               "is the cheapest option." if m["kw_net"] > 0 else
               "At these prices extra KWs are cheaper than add-ons while add-ons are needed.")))

    def fill_whatif(self, s):
        extra = {key: to_number(v.get()) for key, v in self.whatif_vars.items()}
        now = self.result
        after = calculate(self.quantities(extra), s)
        term = max(1, s["term_months"])
        rows = [
            ("EA KW licences", now.ea_licenses, after.ea_licenses, False),
            ("Common Area add-ons", now.ca_addon, after.ca_addon, False),
            ("Access add-ons", now.access_addon, after.access_addon, False),
            ("Monthly cost", now.monthly, after.monthly, True),
            ("Yearly cost", now.monthly * 12, after.monthly * 12, True),
            (f"{term}-month term cost", now.term_total, after.term_total, True),
        ]
        self.whatif.delete(*self.whatif.get_children())
        out = []
        for label, a, b, is_money in rows:
            fmt = money if is_money else (lambda v: f"{v:,}")
            d = b - a
            out.append((label, (fmt(a), fmt(b), ("+" if d > 0 else "") + fmt(d)),
                        (("key",) if label.endswith("term cost") else ()) + (("up",) if d > 0 else ())))
        self._rows(self.whatif, out)
        p_kw, p_ca, p_ac = unit_prices(s)
        self.unit_note.config(text=(
            f"Once the included allowance is used up, each extra Common Area device costs {money(p_ca)}/month "
            f"({money(p_ca * term)} over {term} months) and each extra Access device {money(p_ac)}/month "
            f"({money(p_ac * term)}). Devices still inside the free allowance cost nothing; each new knowledge "
            f"worker costs {money(p_kw)}/month and adds {s['ca_pct']:g}% / {s['access_pct']:g}% free allowance."))

    def clear(self):
        for v in self.qty_vars.values():
            v.set("0")
        self.first_entry.focus_set()

    def save_defaults(self):
        self.settings = self.current_settings()
        save_settings(self.settings)
        messagebox.showinfo(APP_NAME, "Rules saved - they will load next time.")

    def reset_rules(self):
        d = DEFAULT_SETTINGS
        for lic in LICENSE_TYPES:
            self.how_vars[lic].set(TREATMENTS[d["treatment"][lic]])
        for key, _, _ in RULE_FIELDS + PRICE_FIELDS:
            self.rule_vars[key].set(f"{d[key]:g}")
        self.excess_mode.set(d["excess_mode"])

    def copy_summary(self):
        self.clipboard_clear()
        self.clipboard_append(summary_text(self.result, self.current_settings()))
        messagebox.showinfo(APP_NAME, "Summary copied to clipboard.")

    def export_excel(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            messagebox.showerror(APP_NAME, "openpyxl is not installed (py -m pip install openpyxl).")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="FlexSizer - Flex EA sizing.xlsx",
                                            filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        r, s = self.result, self.current_settings()
        wb = Workbook()
        ws = wb.active
        ws.title = "Flex EA sizing"
        head = Font(bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="4F81BD")
        ws.append(["Existing licence", "Current qty", "Sized qty", "Flex user type"])
        for c in ws[1]:
            c.font, c.fill = head, fill
        for lic, qty, sized, how in r.lines:
            ws.append([lic, qty, sized, TREATMENTS[how]])
        ws.append([])
        for key, label, _ in RULE_FIELDS + PRICE_FIELDS:
            ws.append([label, s[key]])
        ws.append(["Devices above included amount",
                   "Common Area / Access add-ons" if s["excess_mode"] == "addon" else "Extra EA KW licences"])
        ws.append([])
        for label, value in self.result_rows(r, s):
            ws.append([label.strip(), value])
            if label in ("Flex EA Knowledge Worker licences", "   Common Area add-on licences",
                         "   Access add-on licences"):
                ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = Font(bold=True)
        ws.append([])
        for label, value in self.cost_rows(r, s):
            ws.append([label, round(value, 2)])
            ws.cell(ws.max_row, 2).number_format = "#,##0.00"
        ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = Font(bold=True)

        sc = wb.create_sheet("More KW vs add-ons")
        sc.append(["Option", "EA KW", "CA add-on", "Access add-on", "Monthly", "Term total", "vs recommended"])
        for c in sc[1]:
            c.font, c.fill = head, fill
        options = kw_scenarios(self.quantities(), s)
        for label, o in options:
            sc.append([label, o.ea_licenses, o.ca_addon, o.access_addon, round(o.monthly, 2),
                       round(o.term_total, 2), round(o.term_total - options[0][1].term_total, 2)])
        for row in sc.iter_rows(min_row=2, min_col=5, max_col=7):
            for c in row:
                c.number_format = "#,##0.00"
        sc.column_dimensions["A"].width = 24
        for col in "BCDEFG":
            sc.column_dimensions[col].width = 16

        ws.append([])
        ws.append(["Source", DATA_SHEET_URL])
        ws.column_dimensions["A"].width = 42
        for col in "BCD":
            ws.column_dimensions[col].width = 20
        try:
            wb.save(path)
        except OSError as e:
            messagebox.showerror(APP_NAME, f"Could not save:\n{e}")
            return
        os.startfile(path)


def setup_style(root):
    s = ttk.Style(root)
    s.theme_use("clam")
    root.configure(background=C["bg"])
    s.configure(".", font=(F, 10), background=C["bg"], foreground=C["ink"])
    # Panels are plain white cards: no inner border, title sits on the card.
    s.configure("TLabelframe", background=C["surface"], borderwidth=0, relief="flat", padding=(16, 6, 16, 14))
    s.configure("TLabelframe.Label", background=C["surface"], foreground=C["navy"], font=(F, 11, "bold"))
    for w in ("TFrame", "TLabel", "TRadiobutton"):
        s.configure(w, background=C["surface"])
    s.map("TRadiobutton", background=[("active", C["surface"])])
    s.configure("App.TFrame", background=C["bg"])
    s.configure("ColHead.TLabel", foreground=C["muted"], font=(F, 9, "bold"))
    s.configure("Muted.TLabel", foreground=C["muted"], font=(F, 9))
    s.configure("Link.TLabel", foreground=C["accent_dark"], font=(F, 9, "underline"))
    s.configure("Foot.TLabel", background=C["bg"], foreground=C["muted"], font=(F, 8))
    s.configure("FootLink.TLabel", background=C["bg"], foreground=C["accent_dark"], font=(F, 8, "underline"))
    s.configure("Vertical.TScrollbar", background=C["bg"], troughcolor=C["surface"], bordercolor=C["surface"],
                arrowcolor=C["muted"], gripcount=0)
    s.configure("TNotebook", background=C["surface"], borderwidth=0, tabmargins=(0, 0, 0, 0))
    s.configure("TNotebook.Tab", font=(F, 10, "bold"), padding=(14, 6), background=C["bg"],
                foreground=C["muted"], bordercolor=C["line"])
    s.map("TNotebook.Tab", background=[("selected", C["accent_soft"])], foreground=[("selected", C["navy"])])
    s.configure("TButton", padding=(12, 6), background=C["surface"], bordercolor=C["line"])
    s.map("TButton", background=[("active", C["accent_soft"])])
    s.configure("Accent.TButton", background=C["accent"], foreground="#FFFFFF", bordercolor=C["accent"],
                font=(F, 10, "bold"))
    s.map("Accent.TButton", background=[("active", C["accent_dark"])])
    s.configure("Ghost.TButton", background=C["navy"], foreground="#FFFFFF", bordercolor="#3D5F82",
                font=(F, 10, "bold"))
    s.map("Ghost.TButton", background=[("active", "#173E66")])
    s.configure("Treeview", rowheight=26, fieldbackground=C["surface"], background=C["surface"], bordercolor=C["line"])
    s.configure("Treeview.Heading", font=(F, 9, "bold"), background=C["bg"], foreground=C["muted"], relief="flat")
    s.map("Treeview", background=[("selected", C["accent_soft"])], foreground=[("selected", C["ink"])])
    for w in ("TEntry", "TSpinbox", "TCombobox"):
        s.configure(w, fieldbackground=C["surface"], bordercolor=C["line"], padding=4)
    s.map("TCombobox", fieldbackground=[("readonly", C["surface"])],
          selectbackground=[("readonly", C["surface"])], selectforeground=[("readonly", C["ink"])])
    root.option_add("*TCombobox*Listbox.font", (F, 10))


def header(root, app):
    bar = tk.Frame(root, background=C["navy"], padx=18, pady=12)
    bar.pack(fill="x")
    tk.Frame(root, background=C["accent"], height=3).pack(fill="x")
    text = tk.Frame(bar, background=C["navy"])
    text.pack(side="left")
    tk.Label(text, text=APP_NAME, bg=C["navy"], fg="#FFFFFF", font=(F, 17, "bold")).pack(anchor="w")
    tk.Label(text, text="Cisco Collaboration Flex Plan 3.0  ·  Enterprise Agreement  ·  On-premises calling",
             bg=C["navy"], fg="#9FC3E0", font=(F, 10)).pack(anchor="w")
    ttk.Button(bar, text="Export to Excel", style="Accent.TButton", command=app.export_excel).pack(side="right")
    ttk.Button(bar, text="Copy summary", style="Ghost.TButton", command=app.copy_summary).pack(side="right", padx=8)


def main():
    if os.name == "nt":
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    root = tk.Tk()
    root.title(f"{APP_NAME} — {APP_TAGLINE}")
    icon = os.path.join(os.path.dirname(os.path.abspath(__file__)), "app.ico")
    if os.path.exists(icon):
        root.iconbitmap(icon)
    setup_style(root)
    app = App(root)
    header(root, app)
    app.pack(fill="both", expand=True)
    scale = app.scale
    root.minsize(int(1180 * scale), int(780 * scale))
    root.mainloop()


if __name__ == "__main__":
    main()
